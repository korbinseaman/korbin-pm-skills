"""Exercise generated report JavaScript with a small DOM and anonymous fixture."""

from __future__ import annotations

import argparse
import importlib.util
import json
import re
import shutil
import subprocess
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/ur-synthesize-report/scripts/generate_report.py"
spec = importlib.util.spec_from_file_location("interactive_report", SCRIPT)
assert spec and spec.loader
report = importlib.util.module_from_spec(spec)
spec.loader.exec_module(report)


def fixture() -> dict:
    questions = [
        {"id": "Q1", "question": "测试样本使用的手机品牌？", "type": "single_choice", "options": ["品牌 A", "品牌 B", "品牌 C"]},
        {"id": "Q2", "question": "测试样本希望优先改进哪些能力？", "type": "multi_choice", "options": ["按相册备份", "管理云端空间", "保护私密相册"]},
        {"id": "Q3", "question": "测试样本对概念的认可程度？", "type": "likert_scale", "options": ["1", "2", "3", "4", "5"]},
        {"id": "Q4", "question": "测试样本的功能优先顺序？", "type": "ranking", "options": ["按相册备份", "管理云端空间"]},
        {"id": "Q5", "question": "测试样本给出的评分？", "type": "nps", "options": []},
    ]
    answers = [
        {"Q1": "品牌 A", "Q2": "按相册备份；保护私密相册", "Q3": 5, "Q4": "按相册备份；管理云端空间", "Q5": 0},
        {"Q1": "品牌 A", "Q2": "管理云端空间", "Q3": 3, "Q4": "管理云端空间；按相册备份", "Q5": 10},
        {"Q1": "品牌 B", "Q2": "按相册备份；管理云端空间", "Q3": 4, "Q4": "按相册备份；管理云端空间", "Q5": 8},
        {"Q1": "品牌 B", "Q3": 2},
    ]
    data = {"questions": questions, "responses": [{"answers": item} for item in answers]}
    return {
        "title": "交互报表功能预览（测试数据）", "topic": "交互报表功能预览（测试数据）", "run_id": "interaction-preview", "generated_at": "2026-09-18T00:00:00",
        "data_source": "synthetic", "sample": {"total": 4, "completed": 4, "analyzable": 4, "failed": 0, "completion_rate": 1, "quality_grade": "excellent", "model_distribution": {}},
        "research_context": {"decision": "验证报表界面的操作与导出", "target_audience": "仅界面测试样本"},
        "data_cleaning": {"completed_before_cleaning": 4, "analyzable_after_cleaning": 4, "excluded_count": 0},
        "interactive_data": data,
        "descriptive_results": [{"question_id": q["id"], "question": q["question"], "type": q["type"], "base_n": sum(q["id"] in a for a in answers), "missing_n": sum(q["id"] not in a for a in answers), "multi_select": q["type"] == "multi_choice", **report.distribution_for(q, [(str(index), a[q["id"]]) for index, a in enumerate(answers) if q["id"] in a])} for q in questions],
        "themes": [], "goal_coverage": [], "cross_tabulations": [], "scale_quality": [], "segment_observations": [], "recommendations": [],
        "report_audit": [
            {"check_id": "statistics", "status": "pending", "note": "4行明确标识的测试数据，仅验证报表功能，不是手机相册调研结论；未与正式调研Excel抽查。"},
            {"check_id": "evidence", "status": "not_applicable", "note": "测试样本无正式定性证据，不补造原话或研究结论。"},
            {"check_id": "charts", "status": "pending", "note": "自动测试覆盖逐题六种视图、独立切换、零频选项、均值/标准差及百分比绘图；50%条宽/柱高为绘图区一半。实际浏览器视觉检查未完成。"},
            {"check_id": "filters", "status": "pending", "note": "自动测试覆盖同题OR、跨题AND、空样本；实际浏览器验收未完成。"},
            {"check_id": "cross", "status": "pending", "note": "自动测试覆盖X两题组合、Y多题、有效分母、缺失及空组。实际浏览器验收未完成。"},
            {"check_id": "tabs", "status": "pending", "note": "四个Tab点击、键盘导航、筛选区隐藏、全批次口径和导出隔离已自动验证，单元测试不替代浏览器验收。"},
            {"check_id": "office", "status": "passed", "note": "测试生成的DOCX/PPTX/XLSX已检查ZIP CRC及全部XML，并用python-docx、python-pptx、openpyxl重开；结论PPT为三页，含百分比与实际分母；Excel选项为文本而非公式。"},
            {"check_id": "desktop_mobile", "status": "pending", "note": "浏览器自动检查受限，实际页面视觉及窄屏检查未完成。"},
            {"check_id": "downloads", "status": "pending", "note": "实际浏览器下载尚未验收；页面不提供保存报告及PNG/SVG导出。"},
        ],
        "limitations": ["本页是验证界面的四行测试数据，不是手机相册课题的调研结果。", "合成模拟数据不代表真实用户或市场总体。"],
    }


HARNESS = r"""
const vm = require('vm'), assert = require('assert');
let input = ''; process.stdin.setEncoding('utf8');
process.stdin.on('data', chunk => input += chunk);
process.stdin.on('end', async () => {
  const {script, office, payload} = JSON.parse(input);
  const downloads = [], blobs = new Map();
  class Element {
    constructor(tag = 'div') { this.nodeType = 1; this.tagName = tag.toUpperCase(); this.children = []; this.attrs = {}; this.listeners = {}; this._html = ''; this._value = ''; this._text = ''; this.options = []; this.classList = {contains: () => false}; }
    set innerHTML(value) { this._html = value; if (this.tagName === 'SELECT') { this.options = [...value.matchAll(/<option value="([^"]*)"([^>]*)>(.*?)<\/option>/g)].map(m => ({value:m[1], selected:m[2].includes('selected'), disabled:false})); this._value = (this.options.find(o => o.selected) || this.options[0] || {}).value || ''; } }
    get innerHTML() { return this._html; }
    set textContent(value) { this._text = value; }
    get textContent() { return this._text; }
    get value() { return this._value; }
    set value(value) { this._value = value; }
    get selectedOptions() { return this.options.filter(o => o.selected); }
    addEventListener(type, handler) { this.listeners[type] = handler; }
    querySelectorAll(selector) {
      if (selector !== "button[data-chart]") return [];
      if (!this.chartButtons || this.buttonHTML !== this._html) {
        this.buttonHTML=this._html;
        this.chartButtons=[...this._html.matchAll(/data-question="([^"]+)" data-chart="([^"]+)"/g)].map(match=>{
          const button=new Element("button");button.dataset={question:match[1],chart:match[2]};return button;
        });
      }
      return this.chartButtons;
    }
    setAttribute(name, value) { this.attrs[name] = value; }
    getAttribute(name) { return this.attrs[name]; }
    hasAttribute(name) { return name in this.attrs; }
    showModal() { this.open = true; }
    close() { this.open = false; }
    appendChild(child) { this.children.push(child); }
    focus() { this.focused = true; }
    click() { if (this.tagName === 'A') downloads.push({name: this.download, blob: blobs.get(this.href)}); else this.listeners.click?.(); }
  }
  const ids = {};
  for (const id of ['report-interactive-data','active-filters','filter-status','view-context','chart-title','chart-note','interactive-charts','distribution-table','cross-title','cross-note','cross-table','add-filter','clear-filters','save-report','download-png','download-svg','export-cross-csv','download-status','ordinary-questions','cross-groups','cross-targets','add-cross-group','add-cross-target','swap-cross','calculate-cross','export-dialog','open-export','close-export','confirm-export','export-error','conclusion-slides','download-conclusion-ppt']) ids[id] = new Element();
  for (const id of ['filter-question','filter-value','chart-question','chart-type','cross-group','cross-target','export-format','export-paper']) ids[id] = new Element('select');
  for (const id of ['tab-ordinary','tab-cross','tab-conclusions','tab-quality','panel-ordinary','panel-cross','panel-conclusions','panel-quality','sample-filters']) ids[id] = new Element();
  ids['filter-options'] = new Element();
  ids['report-interactive-data'].textContent = JSON.stringify(payload);
  const main = new Element('main'), heading = new Element('h1'), paragraph = new Element('p');
  heading.textContent = '验证报告'; paragraph.textContent = '导出内容'; main.children = [heading, paragraph];
  const stateScripts = [], head = {appendChild: element => stateScripts.push(element)};
  const clone = {querySelectorAll: () => [], querySelector: () => head, get outerHTML() { return '<html><head>' + stateScripts.map(s => '<script data-report-state>' + s.textContent + '</script>').join('') + '</head></html>'; }};
  ids['export-format'].value='docx';ids['export-paper'].value='A4';
  const document = {title:'测试报告',querySelectorAll:()=>[],getElementById: id => ids[id], body: {dataset:{runId:'fixture'}}, querySelector: () => main, documentElement:{cloneNode:()=>clone}, createElement: tag => tag === 'canvas' ? {getContext: () => ({measureText: str => ({width: str.length * 15})})} : new Element(tag)};
  const sandbox = {document, window:{}, Blob, TextEncoder, Uint8Array, DataView, URL:{createObjectURL: blob => {const key='blob:' + blobs.size;blobs.set(key,blob);return key;},revokeObjectURL:()=>{}}, setTimeout:()=>{}};
  vm.runInNewContext(office, sandbox);
  vm.runInNewContext(script, sandbox);
  assert(ids['filter-status'].textContent.includes('4 / 4'));
  assert.strictEqual(ids['panel-ordinary'].hidden, false);
  assert.strictEqual(ids['panel-cross'].hidden, true);
  ids['tab-cross'].click();
  assert.strictEqual(ids['panel-ordinary'].hidden, true);
  assert.strictEqual(ids['tab-cross'].getAttribute('aria-selected'), 'true');
  ids['tab-cross'].listeners.keydown({key:'ArrowLeft', preventDefault:()=>{}});
  assert.strictEqual(ids['panel-ordinary'].hidden, false);
  ids['tab-cross'].click();
  assert(ids['cross-table'].innerHTML.includes('品牌 C'), 'empty groups must remain');
  ids['tab-cross'].listeners.keydown({key:'End',preventDefault:()=>{}});
  assert.strictEqual(ids['panel-quality'].hidden,false);
  assert.strictEqual(sandbox.window.ReportAnalysis.officeModel().mode,'quality');
  assert.strictEqual(ids['sample-filters'].hidden,true);
  ids['tab-quality'].listeners.keydown({key:'ArrowLeft',preventDefault:()=>{}});
  assert.strictEqual(ids['panel-conclusions'].hidden,false);
  assert.strictEqual(ids['sample-filters'].hidden,true);
  assert.strictEqual(ids['view-context'].hidden,true);
  assert.strictEqual(ids['panel-cross'].hidden,true);
  ids['tab-quality'].click();
  ids['tab-quality'].listeners.keydown({key:'ArrowRight',preventDefault:()=>{}});
  assert.strictEqual(ids['panel-ordinary'].hidden,false,'keyboard cycles through all four tabs');
  assert.strictEqual(ids['sample-filters'].hidden,false);
  ids['tab-cross'].click();
  const select = (id, value) => { ids[id].value = value; ids[id].listeners.change?.(); };
  select('cross-target', 'Q2');
  const rows = ids['cross-table'].innerHTML;
  assert(rows.includes('<th>品牌 B</th><td>2</td><td>1</td><td>1</td>'), 'missing answers must not enter target denominator');
  assert(rows.includes('100.0%'));
  select('filter-question', 'Q1'); ids['filter-value'].options.forEach(o => o.selected = ['品牌 A','品牌 B'].includes(o.value)); ids['add-filter'].click();
  assert(ids['filter-status'].textContent.includes('4 / 4'), 'same-question choices use OR');
  select('filter-question', 'Q2'); ids['filter-value'].options.forEach(o => o.selected = o.value === '按相册备份'); ids['add-filter'].click();
  assert(ids['filter-status'].textContent.includes('2 / 4'), 'different questions use AND');
  assert(ids['ordinary-questions'].innerHTML.includes('均值 4.50'));
  assert(ids['ordinary-questions'].innerHTML.includes('标准差 0.71'));
  ids['ordinary-questions'].querySelectorAll("button[data-chart]").find(button=>button.dataset.question==='Q1'&&button.dataset.chart==='donut').click();
  assert(ids['ordinary-questions'].innerHTML.split('id="question-Q2"')[0].includes('占比合计'),'per-question mode button works');
  assert(!ids['ordinary-questions'].innerHTML.split('id="question-Q2"')[1].split('id="question-Q3"')[0].includes('<svg'),'other questions retain table view');
  const api=sandbox.window.ReportAnalysis, dist=api.distributionFor(payload.questions[0],payload.responses);
  assert(dist.distribution.some(x=>x.label==='品牌 C'&&x.count===0),'zero-count options remain');
  for(const mode of ['pie','donut','column','bar','line']){
    const svg=api.chartSvg(dist.distribution,mode,'题目');assert(svg.includes('<svg'),'each mode has an SVG');
    assert(svg.includes('%'),'chart values are percentages');
    assert(!svg.includes('NaN')&&!svg.includes('Infinity'));
  }
  assert(api.chartSvg(dist.distribution,'pie','题目').includes('<path'),'pie uses wedges');
  assert(api.chartSvg(dist.distribution,'pie','题目').includes('slice-percent'),'percentages are visible on slices');
  assert(api.chartSvg(dist.distribution,'bar','题目').includes('width="210"'),'50 percent takes half the 100-percent plotting width');
  assert(api.chartSvg(dist.distribution,'column','题目').includes('height="105"'),'50 percent takes half the 100-percent plotting height');
  assert(api.chartSvg(dist.distribution,'donut','题目').includes('占比合计'));
  select('cross-group','Q1');ids['add-cross-group'].click();
  select('cross-group','Q3');ids['add-cross-group'].click();
  select('cross-target','Q2');ids['add-cross-target'].click();
  select('cross-target','Q5');ids['add-cross-target'].click();
  assert.strictEqual(api.crossResults().length,2,'multiple Y targets');
  assert.strictEqual(api.crossResults()[0].rows.length,15,'two X variables use option combinations');
  ids['open-export'].click();assert(ids['export-dialog'].open);
  ids['confirm-export'].click();assert(!ids['export-dialog'].open);
  assert(downloads.at(-1).name.endsWith('.docx'));
  const wordBytes=new Uint8Array(await downloads.at(-1).blob.arrayBuffer());assert(wordBytes[0]===80&&wordBytes[1]===75,'real Office ZIP');
  ids['tab-conclusions'].click();
  const conclusionModel=sandbox.window.ReportAnalysis.officeModel();
  assert.strictEqual(conclusionModel.mode,'conclusions');
  assert.strictEqual(conclusionModel.sections.length,0,'conclusions export does not add filtered question cards');
  assert.strictEqual(conclusionModel.sampleN,4,'conclusions keep full-batch sample');
  assert(conclusionModel.conditions.includes('不随筛选变化'));
  ids['download-conclusion-ppt'].click();
  assert(downloads.at(-1).name.endsWith('.pptx'));
  const conclusionPackage=Buffer.from(await downloads.at(-1).blob.arrayBuffer()).toString('base64');
  ids['tab-cross'].click();
  ids['export-cross-csv'].click(); const csv = await downloads.at(-1).blob.text();
  assert(csv.includes('有效分母') && csv.includes('百分比'));
  ids['clear-filters'].click(); assert(ids['filter-status'].textContent.includes('4 / 4'));
  assert(sandbox.window.ReportAnalysis.distributionFor(payload.questions[4],payload.responses).distribution.some(x=>x.label==='0'&&x.count===1),'zero numeric answers must survive');
  select('filter-question', 'Q1'); ids['filter-value'].options.forEach(o => o.selected = o.value === '品牌 C'); ids['add-filter'].click();
  assert(ids['filter-status'].textContent.includes('0 / 4'));
  assert(!ids['cross-table'].innerHTML.includes('NaN') && !ids['cross-table'].innerHTML.includes('Infinity'));
  const model={title:'离线报表验证',runId:'fixture',conditions:'Q1 = 品牌 A',sampleN:2,mode:'ordinary',paper:'A4',sections:[{kind:'question',title:'Q1 品牌',note:'实际分母 2',items:[{label:'品牌 A',count:2,percent:100},{label:'=安全文本',count:0,percent:0}]},{kind:'cross',title:'Q1 × Q2',note:'实际回答分母',labels:['能力一','能力二'],rows:[{label:'品牌 A',base:2,missing:1,values:[1,2]},{label:'空组',base:0,missing:0,values:[0,0]}]}],narrative:[{title:'研究限制',lines:['仅用于合成数据检查，不代表真实市场。']}]};
  const packages={conclusion_pptx:conclusionPackage};
  for(const format of ['docx','pptx','xlsx'])packages[format]=Buffer.from(await sandbox.window.ReportOffice.build(format,model).arrayBuffer()).toString('base64');
  console.log(JSON.stringify(packages));
});
"""


class InteractiveReportTests(unittest.TestCase):
    def test_generated_javascript_behaviors(self) -> None:
        if not shutil.which("node"):
            self.skipTest("Node.js is required for JavaScript verification")
        analysis = fixture()
        html = report.render_report(analysis)
        self.assertEqual(html.count('id="ordinary-questions"'), 1)
        for removed in ('save-report', 'download-png', 'download-svg'):
            self.assertNotIn('id="' + removed + '"', html)
        from html.parser import HTMLParser
        class PanelParser(HTMLParser):
            def __init__(self):
                super().__init__()
                self.panel = None
                self.conclusion_headings = []
                self.quality_headings = []
                self.heading = False
            def handle_starttag(self, tag, attrs):
                fields = dict(attrs)
                if tag == "article":
                    self.panel = fields.get("id")
                if tag == "h2":
                    self.heading = self.panel in {"panel-conclusions", "panel-quality"}
            def handle_endtag(self, tag):
                if tag == "article":
                    self.panel = None
                if tag == "h2":
                    self.heading = False
            def handle_data(self, text):
                if self.heading:
                    (self.quality_headings if self.panel == "panel-quality" else self.conclusion_headings).append(text)
        parsed = PanelParser()
        parsed.feed(html)
        for removed in ("研究目标", "主题、痛点与需求", "建议与下一步验证"):
            self.assertNotIn(removed, parsed.conclusion_headings)
        self.assertEqual(parsed.quality_headings, ["分群与研究者分析", "全批次描述统计", "数据质量与清理记录", "限制与未回答问题"])
        scripts = re.findall(r"<script>(.*?)</script>", html, re.S)
        result = subprocess.run(["node", "-e", HARNESS], input=json.dumps({"script": scripts[-1], "office": scripts[-2], "payload": {**analysis["interactive_data"], "conclusion_deck": report.conclusion_deck(analysis)}}, ensure_ascii=False), text=True, encoding="utf-8", capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        import base64
        import io
        import zipfile
        import xml.etree.ElementTree as ET
        packages = json.loads(result.stdout)
        for format, encoded in packages.items():
            binary = base64.b64decode(encoded)
            with zipfile.ZipFile(io.BytesIO(binary)) as package:
                self.assertIsNone(package.testzip())
                self.assertIn("[Content_Types].xml", package.namelist())
                for name in package.namelist():
                    if name.endswith((".xml", ".rels")):
                        ET.fromstring(package.read(name))
        try:
            from docx import Document
            from pptx import Presentation
            from openpyxl import load_workbook
        except ImportError:
            return
        doc = Document(io.BytesIO(base64.b64decode(packages["docx"])))
        self.assertGreater(len(doc.tables), 1)
        deck = Presentation(io.BytesIO(base64.b64decode(packages["pptx"])))
        self.assertGreater(len(deck.slides), 2)
        conclusion = Presentation(io.BytesIO(base64.b64decode(packages["conclusion_pptx"])))
        self.assertEqual(len(conclusion.slides), 3)
        slide_text = " ".join(shape.text for slide in conclusion.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn("%", slide_text)
        self.assertIn("实际分母", slide_text)
        workbook = load_workbook(io.BytesIO(base64.b64decode(packages["xlsx"])))
        self.assertGreater(len(workbook.sheetnames), 2)
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", "user labels must not become spreadsheet formulas")

    def test_filter_payload_omits_open_text_and_identity(self) -> None:
        data = report.interactive_data([{"id":"Q1", "type":"single_choice"}, {"id":"Q2", "type":"open_text"}], [{"用户ID":"P001", "姓名":"测试", "Q1":"是", "Q2":"原文", "Q1_回答原因":"原因"}])
        self.assertEqual(data["responses"], [{"answers":{"Q1":"是"}}])
        self.assertEqual([q["id"] for q in data["questions"]], ["Q1"])

    def test_report_filename_uses_date_topic_and_safe_characters(self) -> None:
        self.assertEqual(report.report_filename({"generated_at": "2026-09-18T11:22:33", "topic": "手机相册付费意愿调研"}), "20260918手机相册付费意愿调研_调研报告.html")
        self.assertEqual(report.report_filename({"title": "20260906手机相册问卷报告"}, "20260918"), "20260918手机相册_调研报告.html")
        name = report.report_filename({"topic": '../相册:云/同步?*'}, "20260918")
        self.assertNotRegex(name, r'[<>:"/\\|?*]')
        self.assertEqual(Path(name).name, name)
        for invalid_date in ("20260230", "2026918", "2026-09-18"):
            with self.assertRaises(ValueError):
                report.report_filename({}, invalid_date)

    def test_quality_audit_embeds_failures_and_does_not_certify_pending_checks(self) -> None:
        analysis = fixture()
        analysis["contract_checks"] = {"xlsx_columns": "passed"}
        analysis["report_audit"] = [{"check_id": "statistics", "status": "passed", "note": '<script>alert("x")</script>'}]
        html = report.render_report(analysis, ["HTML 包含外部资源"])
        audit = html.split('id="report-audit"', 1)[1].split('<h2>限制与未回答问题', 1)[0]
        self.assertIn("结构检查失败", audit)
        self.assertIn("HTML 包含外部资源", audit)
        self.assertIn("待检查", audit)
        self.assertIn("&lt;script&gt;", audit)
        self.assertNotIn('<script>alert', audit)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--preview", type=Path)
    args, remainder = parser.parse_known_args()
    if args.preview:
        args.preview.parent.mkdir(parents=True, exist_ok=True)
        analysis = fixture()
        problems = report.machine_quality(report.render_report(analysis), analysis)
        args.preview.write_text(report.render_report(analysis, problems), encoding="utf-8")
        print(args.preview)
    else:
        unittest.main(argv=[__file__, *remainder])
