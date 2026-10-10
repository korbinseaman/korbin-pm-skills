"""Build an editable research handoff PPTX from an audited analysis_summary.json."""
import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from generate_report import conclusion_deck


INK = RGBColor(25, 46, 62)
MUTED = RGBColor(93, 113, 129)
BLUE = RGBColor(9, 151, 240)
PALE = RGBColor(232, 246, 254)


def _text(slide, value, x, y, w, h, size=16, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.text = str(value)
    for paragraph in frame.paragraphs:
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(4)
    return shape


def _rect(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def save_deck(analysis: dict, output: Path) -> Path:
    pages = conclusion_deck(analysis)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for index, page in enumerate(pages, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _rect(slide, 0, 0, 0.12, 7.5, BLUE)
        _text(slide, page["section"], 0.55, 0.30, 5, 0.34, 12, BLUE, True)
        _text(slide, page["title"], 0.55, 0.72, 12.0, 0.65, 27, INK, True)
        _text(slide, f'全批次可分析样本 n={analysis.get("sample", {}).get("analyzable", "—")}', 0.55, 1.46, 6, 0.32, 11, MUTED)
        _rect(slide, 0.55, 1.94, 6.02, 4.59, RGBColor(247, 250, 252))
        is_conclusion = bool(page.get("supporting_data"))
        lines = page["lines"][:7] if is_conclusion else page["lines"][:4]
        for row, line in enumerate(lines):
            step = 0.65 if is_conclusion else 1.06
            _text(slide, f'{row + 1:02d}  {line}', 0.78, 2.10 + row * step,
                  5.52, 0.62 if is_conclusion else 0.99, 11 if is_conclusion else 14, INK)

        chart = page.get("chart") or {}
        _text(slide, chart.get("title") or "数据证据", 6.87, 1.98, 5.75, 0.55, 15, INK, True)
        items = chart.get("items") or []
        if items:
            for row, item in enumerate(items[:6]):
                y = 2.56 + row * 0.55
                pct = max(0, min(100, float(item.get("percent") or 0)))
                _text(slide, str(item.get("label") or "")[:21], 6.88, y, 2.1, 0.37, 11)
                _rect(slide, 9.04, y + 0.04, 2.65, 0.23, PALE)
                if pct:
                    _rect(slide, 9.04, y + 0.04, 2.65 * pct / 100, 0.23, BLUE)
                _text(slide, f'{item.get("count", 0)} · {pct:.1f}%', 11.83, y - 0.01, 1.27, 0.36, 10, MUTED)
        else:
            _text(slide, "本页没有可对应的封闭题图表；证据缺口见左侧。", 6.88, 2.65, 5.7, 0.85, 14, MUTED)
        _text(slide, chart.get("note") or "", 6.88, 6.03, 5.72, 0.49, 10, MUTED)
        _rect(slide, 0.55, 6.72, 12.1, 0.015, PALE)
        _text(slide, f'{page["source"]} · {analysis.get("run_id", "—")} · {page["boundary"]}',
              0.55, 6.88, 11.4, 0.33, 10, MUTED)
        _text(slide, str(index), 12.25, 6.88, 0.3, 0.33, 10, MUTED)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--analysis", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    save_deck(json.loads(args.analysis.read_text(encoding="utf-8")), args.output)
    print(args.output)
